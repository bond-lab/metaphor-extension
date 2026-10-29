"""Tests for scripts/build_animal_sheet.py."""

from __future__ import annotations

import csv
import sys
from pathlib import Path

import pytest

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "scripts"))

from build_animal_sheet import (
    COLUMNS,
    Addition,
    SourceEntry,
    collect_additions,
    language_sheet,
    main,
    master_sheet,
    normalise,
    parse_language,
    previous_summary,
    previous_translation,
    read_prefill,
    write_workbook,
)
from report_animal_metaphors import Row, Tab

DOG_ANIMAL = "omw-en-dog-02084071-n"
DOG_PERSON = "omw-en-dog-10023039-n"
CAT_ANIMAL = "omw-en-cat-02121620-n"
CAT_PERSON = "omw-en-cat-09900153-n"


def make_row(english: str, person_def: str, **kw: str) -> Row:
    """Build a Row with sensible defaults."""
    base = {
        "english": english,
        "language": "",
        "checked": "",
        "comment": "",
        "person_def": person_def,
        "animal_def": f"the animal {english}",
        "animal_sense": "",
        "person_sense": "",
    }
    base.update(kw)
    return Row(**base)


@pytest.fixture
def master() -> list[Row]:
    """A two-animal English master list."""
    return [
        make_row(
            "dog",
            "informal term for a man",
            animal_sense=DOG_ANIMAL,
            person_sense=DOG_PERSON,
            frequency="42",
        ),
        make_row(
            "cat",
            "a spiteful woman gossip",
            animal_sense=CAT_ANIMAL,
            person_sense=CAT_PERSON,
            frequency="7",
        ),
    ]


@pytest.fixture
def tabs(master: list[Row]) -> list[Tab]:
    """Czech tabs for two years plus a free-form Polish list.

    Czech 2024 is an untouched default-Y tab; Czech 2025 judged both pairs and
    added two metaphors (one re-glossing the master dog row).
    """
    index = {r.key: r for r in master}
    czech_2024 = Tab(
        "2024",
        "Czech",
        [
            make_row(
                "dog",
                "informal term for a man",
                language="pes",
                checked="Y",
                animal_sense=DOG_ANIMAL,
                person_sense=DOG_PERSON,
            ),
            make_row(
                "cat",
                "a spiteful woman gossip",
                language="kočka",
                checked="Y",
                animal_sense=CAT_ANIMAL,
                person_sense=CAT_PERSON,
            ),
        ],
        index,
    )
    czech_2025 = Tab(
        "2025",
        "Czech",
        [
            make_row(
                "dog",
                "informal term for a man",
                language="pes, fena",
                checked="N",
                animal_sense=DOG_ANIMAL,
                person_sense=DOG_PERSON,
            ),
            make_row(
                "dog",
                "a loyal person",
                language="pes",
                checked="Y",
                comment="positive",
                animal_sense=DOG_ANIMAL,
                person_sense=DOG_PERSON,
            ),
            make_row(
                "cat",
                "a spiteful woman gossip",
                language="kočka",
                checked="Y",
                animal_sense=CAT_ANIMAL,
                person_sense=CAT_PERSON,
            ),
            make_row("seagull", "a young woman", language="čajka", checked="Y"),
        ],
        index,
    )
    polish = Tab(
        "2024",
        "Polish",
        [
            make_row("pies", "A loyal person", language="pies", checked="Y"),
            make_row("sroka", "A gossip", language="sroka", checked="Y"),
        ],
        index,
        free_form=True,
    )
    return [czech_2024, czech_2025, polish]


def test_normalise_strips_case_and_punctuation() -> None:
    assert normalise("  A Loyal, person! ") == "a loyal person"


def test_collect_additions_merges_by_animal_and_meaning(tabs: list[Tab]) -> None:
    additions = collect_additions(tabs)

    by_key = {a.key: a for a in additions}
    loyal = by_key[("dog", "a loyal person")]
    assert [s.language for s in loyal.sources] == ["Czech", "Polish"]
    assert loyal.animal_sense == DOG_ANIMAL
    assert str(loyal.sources[0]) == "Czech 2025: pes [Y] -- positive"
    assert ("magpie", "a gossip") in by_key, "free-form Polish word is translated"
    assert ("seagull", "a young woman") in by_key


def test_previous_summary_counts_judged_tabs_only(
    master: list[Row], tabs: list[Tab]
) -> None:
    assert previous_summary(tabs, master[0].key) == "N 1"
    assert previous_summary(tabs, master[1].key) == "Y 1"


def test_previous_summary_per_language_flags_default_year(
    master: list[Row], tabs: list[Tab]
) -> None:
    text = previous_summary(tabs, master[0].key, language="Czech")
    assert text == "2025: N; 2024: Y (default)"
    assert previous_summary(tabs, master[0].key, language="Korean") == ""


def test_previous_translation_prefers_latest_year(
    master: list[Row], tabs: list[Tab]
) -> None:
    assert previous_translation(tabs, master[0].key, "Czech") == "pes, fena"


def test_master_sheet_orders_additions_after_their_animal(
    master: list[Row], tabs: list[Tab]
) -> None:
    def resolve(word: str) -> tuple[str, str] | None:
        return {"seagull": ("omw-en-seagull-02041246-n", "a gull")}.get(word)

    rows = master_sheet(master, collect_additions(tabs), tabs, resolve)

    assert [r["English"] for r in rows] == ["dog", "dog", "cat", "magpie", "seagull"]
    assert rows[0]["Source"] == "omw-en"
    assert rows[0]["Previous"] == "N 1"
    assert rows[0]["Frequency"] == "42"
    assert rows[1]["Person Definition"] == "a loyal person"
    assert rows[1]["Animal sense"] == DOG_ANIMAL
    assert rows[1]["Person Sense"] == ""
    assert rows[3]["Animal sense"] == "", "unknown animal stays unresolved"
    assert rows[4]["Animal sense"] == "omw-en-seagull-02041246-n"
    assert rows[4]["Animal Definition"] == "a gull"
    assert all(r["Checked (Y/N)"] == "" for r in rows)
    assert all(set(r) == set(COLUMNS) for r in rows)


def test_read_prefill_indexes_pairs_and_animal_only(tmp_path: Path) -> None:
    path = tmp_path / "prefill.csv"
    with path.open("w", encoding="utf-8", newline="") as fh:
        writer = csv.DictWriter(
            fh,
            fieldnames=[
                "English",
                "Language",
                "Checked (Y/N)",
                "Comment",
                "Animal sense",
                "Person Sense",
                "Animal (wordnet)",
                "Person (wordnet)",
            ],
        )
        writer.writeheader()
        writer.writerow(
            {
                "English": "dog",
                "Language": "犬",
                "Checked (Y/N)": "Y",
                "Comment": "auto: both senses",
                "Animal sense": DOG_ANIMAL,
                "Person Sense": DOG_PERSON,
                "Animal (wordnet)": "犬; イヌ",
                "Person (wordnet)": "犬",
            }
        )

    index = read_prefill(path)

    assert index[(DOG_ANIMAL, DOG_PERSON)]["Language"] == "犬"
    assert index[(DOG_ANIMAL, "")]["Language"] == "犬; イヌ"
    assert read_prefill(None) == {}


def test_language_sheet_fills_from_students_then_prefill(
    master: list[Row], tabs: list[Tab]
) -> None:
    additions = collect_additions(tabs)
    master_rows = master_sheet(master, additions, tabs)
    prefill = {
        (CAT_ANIMAL, CAT_PERSON): {"Language": "kot", "Comment": "auto: both"},
        (DOG_ANIMAL, ""): {"Language": "pies", "Comment": "auto: animal"},
    }

    czech = language_sheet(master_rows, "Czech", tabs, additions, prefill)
    polish = language_sheet(master_rows, "Polish", tabs, additions, prefill)

    assert czech[0]["Language"] == "pes, fena"
    assert czech[0]["Previous"] == "2025: N; 2024: Y (default)"
    assert czech[0]["Comment"] == ""
    assert czech[1]["Language"] == "pes"
    assert czech[1]["Previous"] == "added by Czech 2025 students"
    assert czech[3]["Language"] == "", "no Czech word for the Polish magpie"
    assert polish[1]["Language"] == "pies"
    assert polish[1]["Previous"] == "added by Polish 2024 students"
    assert polish[2]["Language"] == "kot"
    assert polish[2]["Comment"] == "auto: both"
    assert all(r["Checked (Y/N)"] == "" for r in czech + polish)


def test_addition_from_language() -> None:
    addition = Addition(
        "dog", "x", sources=[SourceEntry("Czech", "2025", "pes", "Y", "")]
    )
    assert addition.from_language("Czech").word == "pes"
    assert addition.from_language("Thai") is None


def test_parse_language() -> None:
    assert parse_language("Czech") == ("Czech", None, "")
    assert parse_language("Czech=cs.csv") == ("Czech", Path("cs.csv"), "")
    assert parse_language("Czech=cs.csv=cswn 3.0 + ARASAAC") == (
        "Czech",
        Path("cs.csv"),
        "cswn 3.0 + ARASAAC",
    )


def test_write_workbook_round_trip(tmp_path: Path) -> None:
    from openpyxl import load_workbook

    path = tmp_path / "book.xlsx"
    write_workbook(
        path,
        {
            "English": (["English", "Language"], [{"English": "dog", "Language": ""}]),
            "Languages": (["Language"], [{"Language": "Czech"}]),
        },
    )

    book = load_workbook(path)
    assert book.sheetnames == ["English", "Languages"]
    sheet = book["English"]
    assert [c.value for c in sheet[1]] == ["English", "Language"]
    assert sheet["A2"].value == "dog"
    assert sheet.freeze_panes == "A2"


def write_tab(path: Path, rows: list[list[str]]) -> None:
    """Write a sheet tab CSV with the standard header."""
    header = [
        "English",
        "Language",
        "Checked (Y/N)",
        "Comment",
        "Person Definition",
        "Animal Definition",
        "Animal sense",
        "Person Sense",
        "",
        "Frequency",
    ]
    with path.open("w", encoding="utf-8", newline="") as fh:
        writer = csv.writer(fh)
        writer.writerow(header)
        writer.writerows(rows)


def test_main_builds_all_tabs(tmp_path: Path) -> None:
    data = tmp_path / "data"
    (data / "2025").mkdir(parents=True)
    dog = ["informal term for a man", "a canine", DOG_ANIMAL, DOG_PERSON, "", "42"]
    write_tab(data / "2025" / "English.csv", [["dog", "", "", ""] + dog])
    write_tab(
        data / "2025" / "Czech.csv",
        [
            ["dog", "pes", "N", ""] + dog,
            ["dog", "pes", "Y", ""] + ["a loyal person"] + dog[1:],
        ],
    )
    out = tmp_path / "2026"

    code = main(
        [
            "--year",
            "2026",
            "--data",
            str(data),
            "--out",
            str(out),
            "--no-reference",
            "--language",
            "Czech",
            "--language",
            "Korean==ARASAAC",
        ]
    )

    assert code == 0
    assert sorted(p.name for p in out.iterdir()) == [
        "Czech.csv",
        "English.csv",
        "Korean.csv",
        "Languages.csv",
        "animal-metaphors-2026.xlsx",
    ]
    with (out / "Czech.csv").open(encoding="utf-8", newline="") as fh:
        rows = list(csv.DictReader(fh))
    assert [r["Person Definition"] for r in rows] == [
        "informal term for a man",
        "a loyal person",
    ]
    assert rows[0]["Language"] == "pes"
    assert rows[0]["Previous"] == "2025: N"
    assert rows[1]["Source"] == "Czech 2025: pes [Y]"
    with (out / "Languages.csv").open(encoding="utf-8", newline="") as fh:
        languages = list(csv.DictReader(fh))
    assert languages[1]["Language data from"] == "ARASAAC"
    assert languages[0]["Language data from"] == "2025 students"


def test_main_fails_without_earlier_years(tmp_path: Path) -> None:
    (tmp_path / "2027").mkdir()
    assert main(["--year", "2026", "--data", str(tmp_path), "--no-reference"]) == 1


def test_wordnet_resolver_prefers_frequent_animal_sense() -> None:
    wn = pytest.importorskip("wn")
    if not any(lx.id == "omw-en" and lx.version == "1.4" for lx in wn.lexicons()):
        pytest.skip("omw-en:1.4 not installed")
    from build_animal_sheet import wordnet_resolver

    resolve = wordnet_resolver("omw-en:1.4")

    assert resolve("hen") == ("omw-en-hen-01792640-n", "adult female chicken")
    assert resolve("ant's mouth") is None
